# -*- coding: utf-8 -*-
"""检查导出的 docx 版式（横向、双表、列宽、页眉）。"""
import io
import sys
import urllib.request

from docx import Document
from docx.shared import Cm

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8020"
AID = sys.argv[2] if len(sys.argv) > 2 else "4"

raw = urllib.request.urlopen(f"{BASE}/api/artifact/{AID}/export?fmt=docx", timeout=30).read()
doc = Document(io.BytesIO(raw))
sec = doc.sections[0]
print("bytes", len(raw))
print("orientation", sec.orientation, "page", round(sec.page_width.cm, 2), "x", round(sec.page_height.cm, 2))
print("margins", round(sec.left_margin.cm, 2), round(sec.right_margin.cm, 2))
print("header", repr(sec.header.paragraphs[0].text))
print("footer", repr(sec.footer.paragraphs[0].text))
print("tables", len(doc.tables))
for i, t in enumerate(doc.tables):
    print(f"table{i}: rows={len(t.rows)} cols={len(t.columns)} widths=" +
          ",".join(f"{round(c.width.cm, 2) if c.width else None}" for c in t.columns))
    print("  header:", [c.text[:12] for c in t.rows[0].cells])
    print("  first:", [c.text[:16] for c in t.rows[1].cells])
print("headings", [(p.style.name, p.text[:20]) for p in doc.paragraphs if p.style.name.startswith("Heading")][:6])
