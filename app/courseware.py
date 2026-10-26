# -*- coding: utf-8 -*-
"""互动课件：结构化课件内容 → PPTX / 自包含互动 HTML。

大模型只输出一份结构化 JSON（parse_deck 负责校验与归一化），
PP 与 HTML 都由本地渲染，保证两种资源内容一致、可编辑、无外部依赖。
"""
import html as _html
import io
import json
import re

SLIDE_TYPES = {
    "cover": "封面", "objectives": "学习目标", "concept": "概念讲解",
    "case": "案例分析", "activity": "课堂活动", "quiz": "随堂测验",
    "summary": "总结", "compare": "对比辨析",
}
_INTERACTION_TYPES = {"quiz": "quiz", "测验": "quiz", "测试": "quiz",
                      "poll": "poll", "投票": "poll", "问卷": "poll",
                      "discuss": "discuss", "讨论": "discuss", "思考": "discuss"}

_COLORS = {
    "ink": "123C33", "paper": "F5EFE2", "card": "FFFDF7", "gold": "C9A961",
    "red": "A83A2A", "muted": "6B7770", "line": "D8CEB9",
}


# ================= 解析 / 归一化 =================
def _as_list(value) -> list:
    if isinstance(value, list):
        return [str(x).strip() for x in value if str(x).strip()]
    if isinstance(value, str):
        parts = re.split(r"\n+|；+", value)
        return [p.strip(" -·•\t") for p in parts if p.strip(" -·•\t")]
    return []


def _strip_code(text: str) -> str:
    t = (text or "").strip()
    t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.I)
    t = re.sub(r"\s*```$", "", t)
    return t.strip()


def parse_deck(content) -> dict:
    """解析并归一化；失败抛 ValueError。只做结构整理，不编造内容。"""
    if isinstance(content, dict):
        data = content
    else:
        t = _strip_code(str(content or ""))
        start, end = t.find("{"), t.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("没有找到 JSON 对象")
        try:
            data = json.loads(t[start:end + 1])
        except Exception as e:
            raise ValueError(f"JSON 解析失败：{e}")
    if not isinstance(data, dict):
        raise ValueError("课件 JSON 顶层必须是对象")
    if isinstance(data.get("courseware"), dict):
        data = data["courseware"]
    slides_raw = data.get("slides")
    if not isinstance(slides_raw, list):
        raise ValueError("缺少 slides 数组")
    slides = []
    for raw in slides_raw:
        if not isinstance(raw, dict):
            continue
        st = str(raw.get("type") or raw.get("layout") or "concept").strip().lower()
        if st not in SLIDE_TYPES:
            st = "concept"
        bullets = _as_list(raw.get("bullets") or raw.get("points") or raw.get("items") or raw.get("content"))
        notes = str(raw.get("notes") or raw.get("speaker_notes") or "").strip()
        inter = raw.get("interaction") or {}
        if not isinstance(inter, dict):
            inter = {}
        itype = _INTERACTION_TYPES.get(str(inter.get("type") or "").strip().lower(), "")
        options = _as_list(inter.get("options") or inter.get("choices"))
        answer = str(inter.get("answer") or "").strip()
        explain = str(inter.get("explain") or inter.get("explanation") or "").strip()
        question = str(inter.get("question") or "").strip()
        slide = {
            "type": st,
            "title": str(raw.get("title") or "").strip(),
            "subtitle": str(raw.get("subtitle") or raw.get("kicker") or "").strip(),
            "bullets": bullets[:6],
            "notes": notes,
            "interaction": {
                "type": itype, "question": question, "options": options[:4],
                "answer": answer, "explain": explain,
            } if (itype or question or options) else {},
        }
        slides.append(slide)
    deck = {
        "title": str(data.get("title") or "").strip(),
        "subtitle": str(data.get("subtitle") or "").strip(),
        "duration": str(data.get("duration") or "").strip(),
        "audience": str(data.get("audience") or "").strip(),
        "learning_goals": _as_list(data.get("learning_goals") or data.get("goals")),
        "slides": slides,
    }
    if not deck["slides"]:
        raise ValueError("slides 为空")
    return deck


def _answer_index(answer: str, options: list) -> int:
    """把 B / 2 / 选项原文解析为 0 基下标；无法判断返回 -1。"""
    a = (answer or "").strip()
    if not a or not options:
        return -1
    m = re.match(r"^([A-Da-d])(?:[.、．:：\s]|$)", a)
    if m:
        idx = "ABCD".index(m.group(1).upper())
        return idx if idx < len(options) else -1
    m = re.match(r"^([1-4])(?:[.、．:：\s]|$)", a)
    if m:
        idx = int(m.group(1)) - 1
        return idx if 0 <= idx < len(options) else -1
    for i, opt in enumerate(options):
        if a == opt or a in opt or opt in a:
            return i
    return -1


def validate_deck(content) -> list:
    issues = []
    if not str(content or "").strip():
        return ["内容为空"]
    try:
        deck = parse_deck(content)
    except Exception as e:
        return [f"课件内容不是合法 JSON：{e}"]
    slides = deck["slides"]
    if len(slides) < 6:
        issues.append(f"课件页数过少（{len(slides)} 页），建议 8–14 页")
    elif len(slides) > 18:
        issues.append(f"课件页数过多（{len(slides)} 页），建议 8–14 页")
    if not deck["learning_goals"]:
        issues.append("缺少 learning_goals（学习目标）")
    types = [s["type"] for s in slides]
    if "cover" not in types:
        issues.append("缺少封面页（type=cover）")
    if "summary" not in types:
        issues.append("缺少总结页（type=summary）")
    if "case" not in types:
        issues.append("缺少案例页（type=case）")
    interactive = 0
    for i, s in enumerate(slides, 1):
        if not s["title"]:
            issues.append(f"第{i}页缺少标题")
        if s["type"] not in ("cover", "quiz") and len(s["bullets"]) < 2:
            issues.append(f"第{i}页要点过少（至少 2 条）")
        if not s["notes"] and s["type"] != "cover":
            issues.append(f"第{i}页缺少讲稿备注 notes")
        inter = s.get("interaction") or {}
        if inter:
            interactive += 1
            if not inter.get("question"):
                issues.append(f"第{i}页互动缺少问题")
            if inter.get("type") == "quiz":
                if len(inter.get("options") or []) < 2:
                    issues.append(f"第{i}页测验缺少选项")
                if not inter.get("answer"):
                    issues.append(f"第{i}页测验缺少答案")
    if interactive < 3:
        issues.append(f"互动页过少（{interactive} 页），至少 3 页课堂互动/测验/讨论")
    if len(str(content)) < 800:
        issues.append(f"课件内容偏短（{len(str(content))} 字符，建议 2000+）")
    return issues


# ================= PPTX =================
def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    h = _COLORS.get(hexstr, hexstr)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(run, size=18, color="ink", bold=False, font="微软雅黑"):
    from pptx.util import Pt
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = _rgb(color)


def _textbox(slide, text, x, y, w, h, size=18, color="ink", bold=False,
             align=None, anchor=None, line=1.25):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor or MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align or PP_ALIGN.LEFT
    p.line_spacing = line
    run = p.add_run()
    run.text = text or ""
    _set_font(run, size, color, bold)
    return box


def _rect(slide, x, y, w, h, color, line=None, shape="rect"):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if shape == "round" else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = 1
    else:
        shp.line.fill.background()
    return shp


def _bullets(slide, items, x, y, w, h, size=18, color="ink"):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.line_spacing = 1.18
        r = p.add_run()
        r.text = "•  " + str(item)
        _set_font(r, size, color)
    return box


def _notes(slide, text):
    if not text:
        return
    slide.notes_slide.notes_text_frame.text = text


def _header(slide, course, week, page, total):
    from pptx.util import Inches, Pt
    _rect(slide, 0, 0, Inches(13.333), Inches(7.5), "paper")
    _rect(slide, Inches(0.55), Inches(0.5), Inches(0.08), Inches(6.5), "gold")
    _textbox(slide, f"{course or '课程'} · 第 {week or ''} 周",
             Inches(0.86), Inches(0.36), Inches(9.4), Inches(0.32), 11, "muted")
    _textbox(slide, f"{page:02d} / {total:02d}",
             Inches(11.45), Inches(0.36), Inches(1.3), Inches(0.32), 10, "muted")


def build_pptx(artifact: dict, course=None, week_no=None) -> bytes:
    """生成 16:9、可编辑、带讲稿备注的 PPTX。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    deck = parse_deck(artifact.get("content") or "{}")
    course = course or {}
    week_no = week_no if week_no is not None else artifact.get("week_no")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = deck["slides"]
    total = len(slides)
    title_text = deck["title"] or (course.get("name") or "课堂课件")
    for idx, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        _header(slide, course.get("name") or "", week_no, idx, total)
        st = s["type"]
        bullets = s.get("bullets") or []
        inter = s.get("interaction") or {}
        if st == "cover":
            _textbox(slide, "TEACHING DECK", Inches(0.9), Inches(1.32),
                     Inches(4), Inches(0.35), 12, "gold", True)
            _textbox(slide, title_text, Inches(0.88), Inches(1.72),
                     Inches(8.1), Inches(1.5), 38, "ink", True, line=1.12)
            sub = deck["subtitle"] or s.get("subtitle") or ""
            if sub:
                _textbox(slide, sub, Inches(0.9), Inches(3.35),
                         Inches(7.4), Inches(0.8), 19, "muted", line=1.35)
            if deck["duration"]:
                _textbox(slide, "建议课时：" + deck["duration"], Inches(0.9), Inches(4.22),
                         Inches(6), Inches(0.4), 13, "red")
            goals = deck["learning_goals"][:3]
            if goals:
                _textbox(slide, "学习目标", Inches(9.3), Inches(1.75),
                         Inches(3.2), Inches(0.32), 12, "gold", True)
                _bullets(slide, goals, Inches(9.3), Inches(2.12), Inches(3.2), Inches(3.2),
                         13, "ink")
            _textbox(slide, "案头 · 备课智能体", Inches(0.9), Inches(6.3),
                     Inches(5), Inches(0.3), 10, "muted")
            _notes(slide, s.get("notes") or "开场：问好，说明本课主题与课堂任务。")
            continue

        _textbox(slide, SLIDE_TYPES.get(st, "课堂内容").upper(),
                 Inches(0.9), Inches(1.0), Inches(5), Inches(0.28), 10, "gold", True)
        _textbox(slide, s["title"] or deck["title"], Inches(0.86), Inches(1.28),
                 Inches(11.6), Inches(0.85), 30, "ink", True, line=1.15)
        _rect(slide, Inches(0.9), Inches(2.2), Inches(1.05), Inches(0.05), "gold")

        if st == "quiz":
            q = inter.get("question") or s.get("subtitle") or "请选择正确答案"
            _textbox(slide, q, Inches(0.9), Inches(2.6), Inches(11.2), Inches(0.75),
                     20, "ink", True)
            opts = inter.get("options") or []
            for i, opt in enumerate(opts):
                col, row = i % 2, i // 2
                x = Inches(0.9 + col * 5.65)
                y = Inches(3.55 + row * 1.05)
                _rect(slide, x, y, Inches(5.25), Inches(0.82), "card", "line", "round")
                letter = "ABCD"[i]
                _textbox(slide, letter + ".  " + str(opt), x + Inches(0.18),
                         y + Inches(0.14), Inches(4.9), Inches(0.55), 15, "ink")
            ans = inter.get("answer") or ""
            if ans:
                _textbox(slide, "参考答案：" + ans, Inches(0.9), Inches(5.95),
                         Inches(7), Inches(0.35), 11, "red")
            if inter.get("explain"):
                _textbox(slide, inter["explain"][:110], Inches(0.9), Inches(6.32),
                         Inches(11.2), Inches(0.45), 10, "muted")
        else:
            if bullets:
                size = 23 if len(bullets) <= 3 else (20 if len(bullets) <= 4 else 17)
                _bullets(slide, bullets, Inches(0.9), Inches(2.55),
                         Inches(8.2) if inter else Inches(9.75), Inches(4.15), size)
            if inter:
                box_x = Inches(9.35)
                _rect(slide, box_x, Inches(2.55), Inches(3.1), Inches(3.75), "card", "line", "round")
                _textbox(slide, "课堂互动", box_x + Inches(0.22), Inches(2.75),
                         Inches(2.65), Inches(0.3), 11, "gold", True)
                _textbox(slide, inter.get("question") or "请结合本课内容思考并回答。",
                         box_x + Inches(0.22), Inches(3.12), Inches(2.65), Inches(1.15),
                         14, "ink", True, line=1.3)
                if inter.get("options"):
                    _textbox(slide, " / ".join(inter["options"])[:42],
                             box_x + Inches(0.22), Inches(4.55), Inches(2.65), Inches(0.62),
                             10, "muted")
                if inter.get("answer"):
                    _textbox(slide, "提示：" + inter["answer"][:60],
                             box_x + Inches(0.22), Inches(5.35), Inches(2.65), Inches(0.55),
                             10, "red")
        if not inter:
            _textbox(slide, f"{idx:02d}", Inches(10.72), Inches(4.55),
                     Inches(1.75), Inches(1.72), 86, "E7DCC5", True,
                     align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
        if s.get("subtitle") and st != "cover":
            _textbox(slide, s["subtitle"], Inches(0.9), Inches(6.55),
                     Inches(11.3), Inches(0.3), 10, "muted")
        note = s.get("notes") or ""
        if inter.get("answer") or inter.get("explain"):
            note += "\n\n互动答案：" + str(inter.get("answer") or "")
            if inter.get("explain"):
                note += "\n解析：" + str(inter["explain"])
        _notes(slide, note)
    prs.core_properties.title = title_text
    prs.core_properties.author = "案头 · 备课智能体"
    prs.core_properties.subject = f"第{week_no}周 互动课件"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ================= 互动 HTML =================
def _esc(s) -> str:
    return _html.escape(str(s or ""), quote=True)


def _render_interaction(inter: dict, idx: int) -> str:
    if not inter:
        return ""
    itype = inter.get("type") or "discuss"
    q = _esc(inter.get("question") or "请结合本课内容思考并回答。")
    options = inter.get("options") or []
    answer = inter.get("answer") or ""
    explain = _esc(inter.get("explain") or "")
    correct = _answer_index(answer, options)
    head = {"quiz": "随堂测验", "poll": "现场投票", "discuss": "小组讨论"}.get(itype, "课堂互动")
    out = ['<div class="interact">', '<div class="ih">' + head + '</div>',
           '<div class="iq">' + q + '</div>']
    if options:
        out.append('<div class="opts">')
        for i, opt in enumerate(options):
            mark = " data-correct=\"1\"" if i == correct else " data-correct=\"0\""
            out.append('<button class="opt"' + mark + '><b>' + "ABCD"[i] +
                       '</b><span>' + _esc(opt) + '</span></button>')
        out.append('</div>')
        out.append('<div class="feedback" hidden><b>参考答案：</b>' + _esc(answer) +
                   ('<br><span>' + explain + '</span>' if explain else '') + '</div>')
    else:
        hint = _esc(answer or explain or "先独立思考 1 分钟，再与同桌交换观点。")
        out.append('<button class="reveal">展开讨论提示</button><div class="hintbox" hidden>' +
                   hint + '</div>')
    out.append('</div>')
    return "".join(out)


def build_html(artifact: dict, course=None, week_no=None) -> str:
    deck = parse_deck(artifact.get("content") or "{}")
    course = course or {}
    week_no = week_no if week_no is not None else artifact.get("week_no")
    title = deck["title"] or (course.get("name") or "课堂课件")
    slides_html = []
    total = len(deck["slides"])
    for i, s in enumerate(deck["slides"], 1):
        st = s["type"]
        bullets = "".join("<li>" + _esc(b) + "</li>" for b in (s.get("bullets") or []))
        goals = ""
        if st == "cover" and deck["learning_goals"]:
            goals = '<div class="goals"><b>学习目标</b><ul>' + "".join(
                "<li>" + _esc(g) + "</li>" for g in deck["learning_goals"][:4]) + "</ul></div>"
        inter = _render_interaction(s.get("interaction") or {}, i)
        sub = ('<p class="sub">' + _esc(s.get("subtitle")) + '</p>') if s.get("subtitle") else ""
        meta = ('<div class="duration">建议课时：' + _esc(deck["duration"]) + '</div>') if (st == "cover" and deck["duration"]) else ""
        slides_html.append(
            '<section class="slide ' + st + ('' if i == 1 else '') + '" data-i="' + str(i) + '">' +
            '<div class="kicker">' + _esc(SLIDE_TYPES.get(st, "课堂内容")) + '</div>' +
            '<h1>' + _esc(s["title"] or title) + '</h1>' + sub + meta +
            (('<ul class="points">' + bullets + '</ul>') if bullets else '') +
            goals + inter +
            '<div class="speaker">讲稿：' + _esc(s.get("notes") or "—") + '</div>' +
            '<div class="pg">' + str(i) + ' / ' + str(total) + '</div>' +
            '</section>'
        )
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{_esc(title)}{'' if '课件' in title else ' · 互动课件'}</title>
<style>
:root{{--ink:#123c33;--deep:#0d2c26;--paper:#f5efe2;--card:#fffdf7;--gold:#c9a961;--red:#a83a2a;--muted:#6b7770;--line:#d8ceb9}}
*{{box-sizing:border-box}}
html,body{{margin:0;height:100%;background:var(--paper);color:var(--ink);font-family:"Noto Serif SC","Songti SC",SimSun,serif;overflow:hidden}}
body{{display:flex;flex-direction:column}}
.top{{flex:none;display:flex;align-items:center;gap:12px;padding:12px 22px;background:var(--deep);color:#f3ecd9;font-size:12.5px}}
.top b{{font-family:inherit;font-size:14px;letter-spacing:.5px}}
.top .sp{{flex:1}}
.top button{{border:1px solid rgba(255,255,255,.25);background:rgba(255,255,255,.08);color:#f3ecd9;border-radius:7px;padding:5px 10px;cursor:pointer;font:inherit}}
.top button:hover{{background:rgba(255,255,255,.16)}}
.deck{{flex:1;position:relative;display:flex;align-items:center;justify-content:center;padding:18px 22px 64px}}
.slide{{position:absolute;width:min(1120px,94vw);aspect-ratio:16/9;max-height:calc(100vh - 150px);background:var(--card);border:1px solid var(--line);border-radius:14px;box-shadow:0 18px 50px rgba(35,50,42,.14);padding:clamp(24px,4vw,58px);display:flex;flex-direction:column;opacity:0;transform:translateY(10px);pointer-events:none;transition:opacity .28s,transform .28s;overflow:hidden}}
.slide.active{{opacity:1;transform:none;pointer-events:auto}}
.slide:not(.cover)::after{{content:"";position:absolute;right:42px;bottom:8px;font-family:"Songti SC",SimSun,serif;font-size:clamp(120px,16vw,220px);line-height:1;color:rgba(201,169,97,.10);pointer-events:none;z-index:0}}
.slide.concept::after{{content:"点"}}.slide.case::after{{content:"案"}}.slide.activity::after{{content:"做"}}.slide.quiz::after{{content:"?"}}.slide.summary::after{{content:"结"}}.slide.objectives::after{{content:"标"}}
.slide > *{{position:relative;z-index:1}}
.slide.cover{{background:linear-gradient(135deg,#fffdf7 0%,#f1e8d5 100%)}}
.slide.cover h1{{font-size:clamp(30px,4.6vw,52px);max-width:78%}}
.kicker{{font:600 11px/1.4 system-ui,sans-serif;letter-spacing:2px;color:var(--gold);margin-bottom:14px}}
h1{{font-size:clamp(22px,3.1vw,36px);line-height:1.25;margin:0 0 10px;color:var(--deep);font-weight:700}}
.sub{{color:var(--muted);font-size:15px;margin:0 0 10px;line-height:1.7}}
.duration{{display:inline-block;margin-top:4px;color:var(--red);font-size:12.5px;border-left:3px solid var(--gold);padding-left:9px}}
.points{{margin:18px 0 0;padding:0;list-style:none;display:grid;gap:10px;max-width:92%}}
.points li{{position:relative;padding-left:22px;font-size:clamp(15px,1.65vw,19px);line-height:1.65;color:#27352f}}
.points li:before{{content:"";position:absolute;left:2px;top:.62em;width:7px;height:7px;border-radius:50%;background:var(--gold)}}
.goals{{margin-top:26px;border-top:1px dashed var(--line);padding-top:18px;max-width:70%}}
.goals b{{font-size:12px;color:var(--gold);letter-spacing:1px}}
.goals ul{{margin:8px 0 0;padding-left:18px;color:var(--muted);font-size:13px;line-height:1.8}}
.interact{{margin-top:auto;border:1px solid var(--line);background:#faf5e8;border-radius:11px;padding:15px 18px;max-width:88%}}
.slide.quiz .interact{{margin-top:24px;max-width:100%;background:#fff;padding:22px 26px}}
.ih{{font:600 10px/1.4 system-ui,sans-serif;color:var(--red);letter-spacing:1.5px;margin-bottom:6px}}
.iq{{font-size:16px;line-height:1.6;color:var(--deep);margin-bottom:10px}}
.slide.quiz .iq{{font-size:20px;margin-bottom:16px}}
.opts{{display:grid;grid-template-columns:1fr 1fr;gap:8px}}
.opt{{display:flex;gap:9px;align-items:flex-start;text-align:left;background:#fffdf7;border:1px solid var(--line);border-radius:8px;padding:9px 11px;cursor:pointer;font:inherit;font-size:13px;color:#2a3832;transition:.15s}}
.slide.quiz .opt{{font-size:15px;padding:13px 16px}}
.opt:hover{{border-color:var(--gold);background:#fffaf0}}
.opt b{{color:var(--gold);flex:none}}
.opt.correct{{border-color:#2e7d5b;background:#eef8f1}}
.opt.wrong{{border-color:#c25a47;background:#fdf0ec}}
.feedback{{margin-top:10px;font-size:12px;color:var(--muted);line-height:1.7}}
.feedback b{{color:var(--red)}}
.reveal{{border:1px solid var(--line);background:#fff;border-radius:7px;padding:6px 10px;cursor:pointer;font:inherit;font-size:12px;color:var(--red)}}
.hintbox{{margin-top:8px;font-size:12.5px;color:var(--muted);line-height:1.7}}
.speaker{{display:none;margin-top:auto;border-top:1px dashed var(--line);padding-top:10px;color:var(--muted);font-size:12px;line-height:1.7}}
body.notes .speaker{{display:block}}
body.notes .interact{{margin-top:12px}}
.pg{{position:absolute;right:24px;bottom:16px;font:11px system-ui,sans-serif;color:var(--muted)}}
.ctrl{{position:fixed;left:0;right:0;bottom:0;height:54px;display:flex;align-items:center;gap:12px;padding:0 22px;background:rgba(18,60,51,.96);color:#f3ecd9;z-index:5}}
.ctrl button{{background:rgba(255,255,255,.1);border:1px solid rgba(255,255,255,.2);color:#f3ecd9;border-radius:7px;padding:6px 12px;cursor:pointer;font:inherit;font-size:12px}}
.ctrl button:hover{{background:rgba(255,255,255,.2)}}
.bar{{flex:1;height:4px;background:rgba(255,255,255,.18);border-radius:2px;overflow:hidden}}
.bar i{{display:block;height:100%;width:0;background:var(--gold);transition:width .25s}}
.count{{font:12px system-ui,sans-serif;min-width:52px;text-align:right}}
.help{{position:fixed;inset:0;background:rgba(13,44,38,.72);display:none;place-items:center;z-index:9;color:#f3ecd9;font-size:14px;line-height:2}}
.help.on{{display:grid}}
.help div{{background:#123c33;border:1px solid rgba(255,255,255,.16);border-radius:12px;padding:24px 30px;max-width:520px}}
.help b{{color:var(--gold)}}
@media (max-width:760px){{.slide{{padding:22px;max-height:calc(100vh - 130px)}}.opts{{grid-template-columns:1fr}}.points{{max-width:100%}}.top{{padding:9px 12px;gap:7px;font-size:11px}}.top b{{font-size:12px}}.top .hide-m{{display:none}}.ctrl{{padding:0 12px;gap:7px}}}}
@media print{{html,body{{overflow:visible;background:#fff}}.top,.ctrl,.help,.pg{{display:none}}.deck{{display:block;padding:0}}.slide{{position:relative;opacity:1;transform:none;page-break-after:always;width:100%;max-height:none;box-shadow:none;border:0;border-radius:0}}.speaker{{display:block}}}}
</style>
</head>
<body>
<header class="top">
  <b>{_esc(title)}</b><span class="hide-m">· {_esc(course.get('name') or '')} · 第 {_esc(week_no)} 周</span><span class="sp"></span>
  <button id="btnNotes">讲稿 (N)</button><button id="btnFull">全屏 (F)</button><button id="btnHelp">帮助 (?)</button>
</header>
<main class="deck" id="deck">
{''.join(slides_html)}
</main>
<footer class="ctrl">
  <button id="btnPrev">← 上一页</button><button id="btnNext">下一页 →</button>
  <div class="bar"><i id="bar"></i></div><span class="count" id="count">1 / {total}</span>
</footer>
<div class="help" id="help"><div><b>操作提示</b><br>← / → / 空格：翻页　Home / End：首页 / 末页<br>N：显示 / 隐藏讲稿　F：全屏<br>测验选项点击后显示对错与解析<br>打印：Ctrl / ⌘ + P 可输出全部页面</div></div>
<script>
(function(){{
  var slides=[].slice.call(document.querySelectorAll('.slide')),i=0;
  function show(n){{i=Math.max(0,Math.min(slides.length-1,n));slides.forEach(function(s,k){{s.classList.toggle('active',k===i)}});document.getElementById('bar').style.width=((i+1)/slides.length*100)+'%';document.getElementById('count').textContent=(i+1)+' / '+slides.length;}}
  document.getElementById('btnPrev').onclick=function(){{show(i-1)}};
  document.getElementById('btnNext').onclick=function(){{show(i+1)}};
  document.getElementById('btnNotes').onclick=function(){{document.body.classList.toggle('notes')}};
  document.getElementById('btnFull').onclick=function(){{if(!document.fullscreenElement)document.documentElement.requestFullscreen&&document.documentElement.requestFullscreen();else document.exitFullscreen&&document.exitFullscreen();}};
  document.getElementById('btnHelp').onclick=function(){{document.getElementById('help').classList.toggle('on')}};
  document.getElementById('help').onclick=function(){{this.classList.remove('on')}};
  document.addEventListener('keydown',function(e){{
    if(e.key==='ArrowRight'||e.key===' '||e.key==='PageDown'){{e.preventDefault();show(i+1)}}
    else if(e.key==='ArrowLeft'||e.key==='PageUp'){{e.preventDefault();show(i-1)}}
    else if(e.key==='Home'){{show(0)}}else if(e.key==='End'){{show(slides.length-1)}}
    else if(e.key==='n'||e.key==='N'){{document.body.classList.toggle('notes')}}
    else if(e.key==='f'||e.key==='F'){{document.getElementById('btnFull').click()}}
    else if(e.key==='?'||e.key==='/'){{document.getElementById('help').classList.toggle('on')}}
  }});
  document.querySelectorAll('.opt').forEach(function(btn){{btn.addEventListener('click',function(){{
    var box=btn.closest('.interact');if(!box)return;
    var opts=box.querySelectorAll('.opt');opts.forEach(function(x){{x.classList.remove('correct','wrong')}});
    if(btn.getAttribute('data-correct')==='1')btn.classList.add('correct');else btn.classList.add('wrong');
    var fb=box.querySelector('.feedback');if(fb)fb.hidden=false;
  }})}});
  document.querySelectorAll('.reveal').forEach(function(btn){{btn.addEventListener('click',function(){{
    var h=btn.parentNode.querySelector('.hintbox');if(h)h.hidden=!h.hidden;
  }})}});
  show(0);
}})();
</script>
</body>
</html>"""
