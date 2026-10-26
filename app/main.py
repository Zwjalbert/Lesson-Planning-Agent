# -*- coding: utf-8 -*-
"""FastAPI 入口（v2）：备课智能体 · 本地单机版。

启动：python -m uvicorn app.main:app --port 8000（或 run.bat / run.sh）
前端：http://127.0.0.1:8000

接口分区：
  /api/course*     课程、资料、课程地图、生成任务、看板
  /api/artifact*   产物读取、确认（门 B）、修订、版本、导出
  /api/vault*      公共资源区（全校共享的模板/方法论/案例/思政素材）
  /api/frameworks  提示词框架库（8 要素，只读展示）
  /api/task*       后台任务状态与取消
  /api/usage       用量与费用
"""
import json
import re
import time
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import quote

from fastapi import FastAPI, File, Request, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles

from . import config, store, render, frameworks, rag
from .agent import orchestrator, tools
from .gateway import GatewayError

FILE_CN = render.FILE_CN


@asynccontextmanager
async def lifespan(app: FastAPI):
    store.init()
    yield


app = FastAPI(title="备课智能体", version=config.APP_VERSION, lifespan=lifespan)
app.mount("/static", StaticFiles(directory=str(config.STATIC_DIR)), name="static")


# ================= 通用 =================
def ok(**kw):
    return {"ok": True, **kw}


def fail(msg, status=200):
    return JSONResponse({"ok": False, "msg": str(msg)}, status_code=status)


@app.exception_handler(GatewayError)
async def _gateway_error(request: Request, exc: GatewayError):
    return JSONResponse({"ok": False, "msg": str(exc)}, status_code=200)


@app.exception_handler(Exception)
async def _any_error(request: Request, exc: Exception):
    return JSONResponse({"ok": False, "msg": f"服务异常：{type(exc).__name__}: {str(exc)[:300]}"},
                        status_code=500)


async def payload(request: Request) -> dict:
    """同时接受 JSON 与 form/multipart，兼容旧调用方式。"""
    ct = (request.headers.get("content-type") or "").lower()
    if "application/json" in ct:
        try:
            data = await request.json()
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}
    try:
        form = await request.form()
        return {k: v for k, v in form.items()}
    except Exception:
        return {}


def as_int(v, default=0):
    try:
        return int(str(v).strip())
    except Exception:
        return default


def as_list(v) -> list:
    """'2,3' / '[2,3]' / '2-4' / list → [2,3]"""
    if v is None or v == "":
        return []
    if isinstance(v, (list, tuple)):
        raw = list(v)
    else:
        s = str(v).strip()
        if s.startswith("["):
            try:
                raw = json.loads(s)
            except Exception:
                raw = re.split(r"[,，、\s]+", s.strip("[]"))
        else:
            raw = re.split(r"[,，、\s]+", s)
    out = []
    for x in raw:
        if isinstance(x, str) and re.match(r"^\d+\s*[-~—]\s*\d+$", x):
            a, b = re.split(r"[-~—]", x)
            out.extend(range(as_int(a), as_int(b) + 1))
        else:
            n = as_int(x)
            if n:
                out.append(n)
    return sorted(set(out))


def as_bool(v, default=False) -> bool:
    if v is None:
        return default
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in ("1", "true", "yes", "on")


def _course(cid):
    return store.get_course(cid)


def _artifact_or_course(aid, cid=None):
    a = store.get_artifact(aid)
    if not a or (cid is not None and a["course_id"] != int(cid)):
        return None
    return a


def _content_disposition(filename: str) -> str:
    return f"attachment; filename*=UTF-8''{quote(filename)}"


# ================= 系统 =================
@app.get("/api/system/info")
def system_info():
    return ok(**config.info(),
              frameworks=frameworks.counts(),
              vault=store.vault_summary(),
              skills=[{"name": s["name"], "gate": s["gate"], "desc": s["desc"]}
                      for s in orchestrator.SKILLS])


@app.get("/api/health")
def health():
    try:
        store.init()
        return ok(db=str(config.DB_PATH), output=str(config.OUTPUT_DIR), mode=config.MODE)
    except Exception as e:
        return fail(f"数据库不可用：{e}", 500)


# ================= 课程 =================
@app.post("/api/course")
async def create_course(request: Request):
    d = await payload(request)
    name = str(d.get("name") or "").strip()
    if not name:
        return fail("请填写课程名称", 400)
    weeks = as_int(d.get("weeks"), 16)
    if not 1 <= weeks <= 30:
        return fail("总周数需在 1–30 之间", 400)
    cid = store.create_course(
        name[:80], str(d.get("semester") or "").strip()[:40], weeks,
        str(d.get("theme") or "").strip()[:200],
        str(d.get("class_name") or "").strip()[:60],
        str(d.get("textbook") or "").strip()[:120])
    store.add_confirm_log(cid, None, "create", name)
    return ok(course_id=cid, course=store.get_course(cid))


@app.get("/api/courses")
def courses():
    return ok(courses=store.list_courses())


@app.get("/api/course/{cid}")
def get_course(cid: int):
    c = _course(cid)
    return ok(course=c) if c else fail("课程不存在", 404)


@app.patch("/api/course/{cid}")
async def update_course(cid: int, request: Request):
    c = _course(cid)
    if not c:
        return fail("课程不存在", 404)
    d = await payload(request)
    fields = {}
    for k in ("name", "semester", "theme", "class_name", "textbook", "status"):
        if k in d:
            fields[k] = str(d[k]).strip()
    if "weeks" in d:
        w = as_int(d["weeks"], c["weeks"])
        if not 1 <= w <= 30:
            return fail("总周数需在 1–30 之间", 400)
        fields["weeks"] = w
    if "auto_chain" in d:
        fields["auto_chain"] = 1 if as_bool(d["auto_chain"], True) else 0
    if "auto_types" in d:
        fts = [f for f in as_list_types(d["auto_types"]) if f in tools.FILE_TYPES]
        fields["auto_types"] = fts or ["lesson_plan", "handout", "courseware", "script"]
    store.update_course(cid, **fields)
    return ok(course=store.get_course(cid))


@app.delete("/api/course/{cid}")
def delete_course(cid: int):
    if not _course(cid):
        return fail("课程不存在", 404)
    store.delete_course(cid)
    return ok(msg="课程及数据已删除")


def as_list_types(v) -> list:
    if v is None:
        return []
    if isinstance(v, list):
        return [str(x) for x in v]
    s = str(v).strip()
    if s.startswith("["):
        try:
            return [str(x) for x in json.loads(s)]
        except Exception:
            pass
    return [x for x in re.split(r"[,，、\s]+", s) if x]


@app.get("/api/course/{cid}/board")
def board(cid: int):
    c = _course(cid)
    if not c:
        return fail("课程不存在", 404)
    plans = store.list_week_plans(cid)
    arts = store.list_artifacts(cid, with_content=False)
    by = {}
    for a in arts:
        by.setdefault(a["week_no"], {})[a["file_type"]] = a["status"]
    for p in plans:
        p["artifact_status"] = by.get(p["week_no"], {})
    tasks = store.list_tasks(cid, 8)
    for t in tasks:
        t.pop("log", None)
    return ok(
        course=c, materials=store.list_materials(cid), week_plans=plans, artifacts=arts,
        usage=store.usage_summary(cid), tasks=tasks, mode=config.MODE,
        messages=[{"role": m["role"], "text": m["text"], "ts": m["ts"]}
                  for m in store.recent_messages(cid, 20)],
        confirm_logs=store.list_confirm_logs(cid, 20),
        rag=rag.stats(cid),
        file_types=[{"id": k, "cn": v["cn"], "stage": v["stage"]}
                    for k, v in sorted(tools.FILE_TYPES.items(), key=lambda x: x[1]["stage"])],
    )


# ================= 资料 =================
@app.post("/api/course/{cid}/material")
async def upload_material(cid: int, file: UploadFile = File(...)):
    if not _course(cid):
        return fail("课程不存在", 404)
    name = file.filename or "未命名"
    ext = Path(name).suffix.lower()
    limit = config.MAX_UPLOAD_MB * 1024 * 1024
    if getattr(file, "size", None) and file.size > limit:
        return fail(f"文件超过 {config.MAX_UPLOAD_MB}MB 限制，请拆分后上传", 400)
    raw = await file.read()
    if len(raw) > limit:
        return fail(f"文件超过 {config.MAX_UPLOAD_MB}MB 限制，请拆分后上传", 400)
    hint = _check_magic(ext, raw)
    if hint:
        return fail(hint, 400)
    try:
        text = await run_in_threadpool(_extract_text, ext, raw)
    except Exception as e:
        return fail(_friendly_err(ext, e), 400)
    if not (text or "").strip():
        return fail("文件中没有可提取的文字", 400)
    mid, replaced = store.add_material(cid, name, text, "")
    store.add_confirm_log(cid, None, "upload", name)
    n = len(text)
    msg = f"已{'更新' if replaced else '上传'}《{name}》（{n} 字符）"
    if n > config.MAX_MATERIAL_CHARS:
        msg += f"，超出单份上限，已截断至 {config.MAX_MATERIAL_CHARS} 字符"
    return ok(msg=msg, material_id=mid, chars=n)


def _check_magic(ext: str, raw: bytes) -> str:
    head = raw[:8]
    if ext in (".docx", ".xlsx"):
        if head[:2] == b"PK":
            return ""
        if head[:4] == b"\xd0\xcf\x11\xe0":
            return (f"这是旧版 Office 二进制格式（.doc/.xls），并非真正的 {ext}。"
                    "请在 Word/Excel 中「另存为」.docx/.xlsx 后再上传。")
        return f"文件不是有效的 {ext}（可能被改过扩展名、损坏或传输不完整）。"
    if ext in (".doc", ".xls"):
        return f"暂不支持旧版 {ext} 二进制格式：请在 Office 中另存为 .docx/.xlsx 后上传，或直接粘贴文本。"
    if ext == ".pdf" and head[:4] != b"%PDF":
        return "文件不是有效的 PDF（可能被改过扩展名或已损坏）。"
    return ""


def _friendly_err(ext: str, e: Exception) -> str:
    m = str(e).lower()
    if ext == ".pdf" and ("未提取到" in str(e) or "ocr" in m or "扫描" in str(e)):
        return str(e)
    if "zip" in m or "ole" in m or "central directory" in m:
        return ("文件不是有效的 Word/Excel 文档（可能被改过扩展名、损坏或传输不完整）。"
                "请在 Office 中另存为新格式后再上传。")
    if "password" in m or "encrypted" in m:
        return "文件已加密，请先解除密码保护后再上传。"
    if ext in (".doc", ".xls"):
        return f"暂不支持旧版 {ext} 格式，请另存为 .docx/.xlsx 后上传。"
    return f"解析失败：{e}"


def _extract_text(ext: str, raw: bytes) -> str:
    if ext in (".txt", ".md", ".markdown", ".csv"):
        text = ""
        for enc in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            text = raw.decode("utf-8", errors="ignore")
        if ext == ".csv":
            import csv
            import io
            try:
                rows = list(csv.reader(io.StringIO(text)))
                text = "\n".join("\t".join(c.strip() for c in row) for row in rows if any(row))
            except Exception:
                pass
        return text
    if ext == ".docx":
        from docx import Document
        import io
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for ti, tb in enumerate(doc.tables):
            parts.append(f"【文档表格 {ti + 1}】")
            for row in tb.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                line = " | ".join(cells).strip(" |")
                if line:
                    parts.append(line)
        return "\n".join(parts)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"【工作表 {ws.title}】")
            for row in ws.iter_rows(values_only=True):
                line = "\t".join("" if c is None else str(c) for c in row).strip("\t")
                if line:
                    parts.append(line)
        try:
            wb.close()
        except Exception:
            pass
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("文件已加密，请先解除密码保护后再上传。")
        texts = [(p.extract_text() or "").strip() for p in reader.pages]
        joined = "\n".join(t for t in texts if t)
        if not joined.strip():
            raise ValueError("PDF 未提取到文字：可能是扫描件/图片型 PDF。演示版暂不支持 OCR，"
                             "请在 Word 中另存为 .docx，或使用文本型 PDF。")
        return joined
    raise ValueError("仅支持 txt / md / csv / docx / xlsx / pdf")


@app.post("/api/course/{cid}/material-text")
async def upload_text(cid: int, request: Request):
    if not _course(cid):
        return fail("课程不存在", 404)
    d = await payload(request)
    filename = str(d.get("filename") or f"资料{int(time.time())}").strip()[:120]
    content = str(d.get("content") or "")
    if not content.strip():
        return fail("内容为空", 400)
    mid, replaced = store.add_material(cid, filename, content, "")
    store.add_confirm_log(cid, None, "upload_text", filename)
    return ok(msg=f"已{'更新' if replaced else '保存'}《{filename}》（{len(content)} 字符）",
              material_id=mid, chars=len(content))


@app.get("/api/course/{cid}/material/{mid}")
def get_material(cid: int, mid: int):
    m = store.get_material(mid)
    if not m or m["course_id"] != cid:
        return fail("资料不存在", 404)
    return ok(material=m)


@app.delete("/api/course/{cid}/material/{mid}")
def delete_material(cid: int, mid: int):
    m = store.get_material(mid)
    if not m or m["course_id"] != cid:
        return fail("资料不存在", 404)
    store.delete_material(mid)
    store.add_confirm_log(cid, None, "delete_material", m["filename"])
    return ok(msg="资料已删除")


@app.post("/api/course/{cid}/parse")
async def parse_materials(cid: int, request: Request):
    if not _course(cid):
        return fail("课程不存在", 404)
    d = await payload(request)
    r = await run_in_threadpool(tools.parse_materials, cid, as_bool(d.get("force"), False))
    return r


# ================= 课程地图（门 A） =================
@app.post("/api/course/{cid}/plan")
async def build_plan(cid: int, request: Request):
    if not _course(cid):
        return fail("课程不存在", 404)
    d = await payload(request)
    weeks = as_list(d.get("weeks"))
    c = _course(cid)
    if not weeks:
        weeks = list(range(1, c["weeks"] + 1))
    return await run_in_threadpool(tools.build_plan, cid, weeks,
                                   str(d.get("focus") or ""), str(d.get("tone") or ""))


@app.patch("/api/course/{cid}/plan/{week_no}")
async def patch_week_plan(cid: int, week_no: int, request: Request):
    d = await payload(request)
    fields = {}
    for k in ("title", "objective", "ideo", "status"):
        if k in d:
            fields[k] = str(d[k]).strip()
    if "hours" in d:
        h = as_int(d["hours"], 2)
        if h not in (2, 3, 4):
            return fail("学时应为 2/3/4", 400)
        fields["hours"] = h
    cur = store.get_week_plan(cid, week_no)
    if not cur:
        return fail("周次不存在", 404)
    # 已批准的周次被修改内容后自动退回待确认，避免「地图已改、产物还是旧的」
    content_keys = ("title", "hours", "objective", "ideo")
    changed = any(k in fields and str(fields[k]) != str(cur.get(k)) for k in content_keys)
    if cur["status"] == "approved" and "status" not in fields and changed:
        fields["status"] = "review"
    return tools.update_week_plan(cid, week_no, **fields)


@app.post("/api/course/{cid}/plan/{week_no}/regenerate")
async def regenerate_week(cid: int, week_no: int, request: Request):
    d = await payload(request)
    return await run_in_threadpool(tools.regenerate_week, cid, week_no,
                                   str(d.get("focus") or ""))


@app.post("/api/course/{cid}/plan/approve")
async def approve_plan(cid: int, request: Request):
    if not _course(cid):
        return fail("课程不存在", 404)
    d = await payload(request)
    weeks = as_list(d.get("weeks")) or None
    return tools.approve_plan(cid, weeks)


# ================= 生成任务（门 B） =================
@app.post("/api/course/{cid}/generate")
async def generate(cid: int, request: Request):
    if not _course(cid):
        return fail("课程不存在", 404)
    d = await payload(request)
    weeks = as_list(d.get("weeks"))
    fts = as_list_types(d.get("file_types")) or None
    return orchestrator.generate_weeks(cid, weeks, fts,
                                       str(d.get("focus") or ""), str(d.get("tone") or ""))


@app.get("/api/task/{tid}")
def task_state(tid: str):
    return ok(**orchestrator.task_state(tid))


@app.post("/api/task/{tid}/cancel")
def cancel_task(tid: str):
    return orchestrator.cancel(tid)


@app.get("/api/course/{cid}/tasks")
def course_tasks(cid: int):
    return ok(tasks=store.list_tasks(cid, 10))


# ================= 产物 =================
@app.get("/api/artifact/{aid}")
def get_artifact(aid: int):
    a = store.get_artifact(aid)
    if not a:
        return fail("产物不存在", 404)
    a["file_type_cn"] = FILE_CN.get(a["file_type"], a["file_type"])
    a["issues"] = tools.validate_text(a["file_type"], a.get("content") or "",
                                     _week_hours(a))
    a["revisions"] = store.list_revisions(aid, with_content=False)
    return ok(artifact=a)


def _week_hours(a: dict) -> int:
    w = store.get_week_plan(a["course_id"], a["week_no"]) or {}
    return int(w.get("hours") or 2)


@app.get("/api/course/{cid}/artifact/{aid}")
def get_course_artifact(cid: int, aid: int):
    a = _artifact_or_course(aid, cid)
    if not a:
        return fail("产物不存在", 404)
    a["file_type_cn"] = FILE_CN.get(a["file_type"], a["file_type"])
    a["issues"] = tools.validate_text(a["file_type"], a.get("content") or "", _week_hours(a))
    a["revisions"] = store.list_revisions(aid, with_content=False)
    return ok(artifact=a)


@app.post("/api/artifact/{aid}/approve")
def approve_artifact(aid: int):
    r = tools.approve_artifact(aid)
    if r.get("ok"):
        a = r.get("artifact") or store.get_artifact(aid)
        auto = orchestrator.on_artifact_approved(a["course_id"], a["week_no"], a["file_type"])
        r["auto"] = auto
    return r


@app.post("/api/artifact/{aid}/revise")
async def revise_artifact(aid: int, request: Request):
    d = await payload(request)
    fb = str(d.get("feedback") or "").strip()
    if not fb:
        return fail("请填写修改要求", 400)
    return await run_in_threadpool(tools.revise_artifact, aid, fb, str(d.get("focus") or ""))


@app.post("/api/artifact/{aid}/reopen")
async def reopen_artifact(aid: int, request: Request):
    d = await payload(request)
    return tools.reopen_artifact(aid, str(d.get("reason") or "退回修改"))


@app.post("/api/artifact/{aid}/reset")
def reset_artifact(aid: int):
    return tools.reset_artifact(aid)


@app.get("/api/artifact/{aid}/revisions")
def revisions(aid: int):
    return ok(revisions=store.list_revisions(aid, with_content=False))


@app.get("/api/artifact/{aid}/revision/{rid}")
def revision(aid: int, rid: int):
    r = store.get_revision(rid)
    if not r or r["artifact_id"] != aid:
        return fail("版本不存在", 404)
    return ok(revision=r)


@app.post("/api/artifact/{aid}/restore/{rid}")
def restore_revision(aid: int, rid: int):
    a = store.get_artifact(aid)
    r = store.get_revision(rid)
    if not a or not r or r["artifact_id"] != aid:
        return fail("版本不存在", 404)
    store.set_artifact(aid, status="review", content=r["content"], warnings="",
                       word_count=len(r["content"]), feedback="", revision=None)
    store.save_revision(aid, r["content"], feedback=f"回退到 v{r['version']}", source="restore",
                        model=r.get("model", ""))
    store.add_confirm_log(a["course_id"], aid, "restore", f"v{r['version']}")
    return ok(msg=f"已回退到第 {r['version']} 版", content=r["content"])


@app.get("/api/artifact/{aid}/export")
def export_artifact(aid: int, fmt: str = "docx"):
    a = store.get_artifact(aid)
    if not a:
        return fail("产物不存在", 404)
    fmt = (fmt or "docx").lower()
    allowed = ("pptx", "html") if a["file_type"] == "courseware" else ("docx", "html", "md")
    if fmt not in allowed:
        return fail("课件仅支持 pptx/html；其他文件支持 docx/html/md", 400)
    data, media, filename = render.export_bytes(aid, fmt)
    return Response(data, media_type=media,
                    headers={"Content-Disposition": _content_disposition(filename)})


@app.get("/api/course/{cid}/artifact/{aid}/export")
def export_course_artifact(cid: int, aid: int, fmt: str = "docx"):
    if not _artifact_or_course(aid, cid):
        return fail("产物不存在", 404)
    return export_artifact(aid, fmt)


@app.get("/api/artifact/{aid}/preview", response_class=HTMLResponse)
def preview_artifact(aid: int):
    """课件在查看页内嵌预览；其他产物回退为普通文章 HTML。"""
    a = store.get_artifact(aid)
    if not a:
        return HTMLResponse("<h3 style='font-family:sans-serif;padding:40px'>产物不存在</h3>",
                            status_code=404)
    if a["file_type"] == "courseware":
        return HTMLResponse(render.courseware_html(a))
    return HTMLResponse(render.artifact_html(a))


@app.post("/api/course/{cid}/artifact/{aid}/edit")
async def artifact_edit(cid: int, aid: int, request: Request):
    a = _artifact_or_course(aid, cid)
    if not a:
        return fail("产物不存在", 404)
    d = await payload(request)
    requirement = str(d.get("requirement") or d.get("feedback") or "").strip()
    if not requirement:
        return fail("请填写修改要求", 400)
    r = await run_in_threadpool(tools.revise_artifact, aid, requirement)
    if not r.get("ok"):
        return r
    a2 = store.get_artifact(aid)
    return ok(content=a2.get("content") or "", model=a2.get("model", ""),
              warns="；".join(tools.validate_text(a2["file_type"], a2.get("content") or "",
                                                  _week_hours(a2))),
              artifact_id=aid, msg=r.get("msg"))


# ================= 资料查看页 =================
@app.get("/api/course/{cid}/artifact/{aid}/viewer", response_class=HTMLResponse)
def artifact_viewer(cid: int, aid: int):
    """逐文件查看页：左侧完整渲染、右侧对 AI 提修改要求（门 B 修改器）。"""
    a = _artifact_or_course(aid, cid)
    if not a:
        return HTMLResponse("<h3 style='font-family:sans-serif;padding:40px'>产物不存在</h3>",
                            status_code=404)
    return HTMLResponse((config.STATIC_DIR / "viewer.html").read_text(encoding="utf-8"))


@app.get("/viewer/{cid}/{aid}", response_class=HTMLResponse)
def viewer_friendly(cid: int, aid: int):
    return artifact_viewer(cid, aid)


# ================= 对话总控 =================
@app.post("/api/course/{cid}/chat")
async def chat(cid: int, request: Request):
    if not _course(cid):
        return fail("课程不存在", 404)
    d = await payload(request)
    text = str(d.get("text") or "").strip()
    if not text:
        return fail("请输入需求", 400)
    return await run_in_threadpool(orchestrator.chat, cid, text, as_bool(d.get("execute"), True))


@app.get("/api/course/{cid}/messages")
def messages(cid: int):
    return ok(messages=[{"role": m["role"], "text": m["text"], "ts": m["ts"]}
                        for m in store.recent_messages(cid, 40)])


@app.delete("/api/course/{cid}/messages")
def clear_messages(cid: int):
    store.clear_messages(cid)
    return ok(msg="对话已清空")


# ================= 公共资源区 =================
@app.get("/api/vault")
def list_vault(kind: str = ""):
    return ok(items=store.list_vault(kind or None), kinds=store.VAULT_KINDS)


@app.post("/api/vault")
async def create_vault(request: Request):
    d = await payload(request)
    name = str(d.get("name") or "").strip()
    content = str(d.get("content") or "").strip()
    if not name or not content:
        return fail("名称与内容不能为空", 400)
    item_id = store.create_vault(str(d.get("kind") or "other"), name[:120], content,
                                 str(d.get("tags") or "")[:200])
    return ok(item_id=item_id, msg="已添加")


@app.patch("/api/vault/{item_id}")
async def update_vault(item_id: int, request: Request):
    if not store.get_vault(item_id):
        return fail("条目不存在", 404)
    d = await payload(request)
    fields = {}
    for k in ("kind", "name", "content", "tags"):
        if k in d:
            fields[k] = str(d[k])
    if "enabled" in d:
        fields["enabled"] = 1 if as_bool(d["enabled"], True) else 0
    store.update_vault(item_id, **fields)
    return ok(msg="已保存")


@app.delete("/api/vault/{item_id}")
def delete_vault(item_id: int):
    if not store.get_vault(item_id):
        return fail("条目不存在", 404)
    store.delete_vault(item_id)
    return ok(msg="已删除")


# ================= 提示词框架 =================
@app.get("/api/frameworks")
def list_frameworks():
    return ok(frameworks=frameworks.all_frameworks())


@app.get("/api/frameworks/{fid}")
def get_framework(fid: str):
    f = frameworks.public_detail(fid)
    return ok(framework=f) if f else fail("框架不存在", 404)


# ================= 用量 / 交付（门 C） =================
@app.get("/api/course/{cid}/usage")
def usage(cid: int):
    return ok(**store.usage_summary(cid))


@app.get("/api/usage")
def usage_all():
    return ok(**store.usage_summary(None))


@app.post("/api/course/{cid}/package")
def package(cid: int):
    try:
        path = render.package_course(cid)
    except Exception as e:
        return fail(str(e), 400)
    store.update_course_status(cid, "delivered")
    p = Path(path)
    return ok(path=str(p), filename=p.name, size=p.stat().st_size if p.exists() else 0)


@app.get("/api/course/{cid}/download")
def download_package(cid: int):
    path = render.latest_package(cid)
    if not path:
        return fail("尚未打包，或历史包已被移动/删除", 404)
    p = Path(path)
    return FileResponse(str(p), filename=p.name,
                        media_type="application/zip")


@app.get("/api/course/{cid}/packages")
def packages(cid: int):
    out = []
    for log in store.list_confirm_logs(cid, 100):
        if log["action"] != "package":
            continue
        p = Path(log["reason"] or "")
        if p.is_file():
            out.append({"path": str(p), "filename": p.name,
                        "size": p.stat().st_size, "ts": log["ts"]})
    return ok(packages=out[:10])


# ================= 首页 =================
@app.get("/")
def index():
    return FileResponse(str(config.STATIC_DIR / "index.html"),
                        headers={"Cache-Control": "no-store"})
